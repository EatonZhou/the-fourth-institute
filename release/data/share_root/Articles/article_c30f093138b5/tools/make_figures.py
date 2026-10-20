"""Generate editable PowerPoint and vector PDF figures from one scene description.
All boxes, labels and connectors are native PowerPoint objects; no bitmap slides.
"""
from pathlib import Path
import json, math, html, unicodedata
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
import cairosvg

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'figures'
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(12), Inches(6)
FONT='Noto Sans CJK SC'
INK='263442'; BLUE='275D85'; TEAL='26756C'; AMBER='99651B'
PALE='EFF4F8'; GREEN='EDF5F1'; GOLD='FAF3E7'; GRAY='F3F4F5'; LINE='91A2AE'
figures=[]

class Fig:
    def __init__(self, name, title):
        self.name=name; self.items=[]; self.slide=prs.slides.add_slide(prs.slide_layouts[6])
        self.txt(28,15,1144,37,title,26,True,BLUE,'left')
        self.line([(28,62),(1172,62)],False,LINE)
        self.txt(30,568,1140,22,'设计示意 · 实线：工作或依赖关系；虚线：反馈或关联；不表示实验测量',13,False,'61717E','left')
        figures.append(self)
    def txt(self,x,y,w,h,text,size=19,bold=False,color=INK,align='center'):
        def wrap(s, fs):
            result=[]
            for raw in s.split('\n'):
                line=''; width=0
                for ch in raw:
                    cw=fs*(1 if unicodedata.east_asian_width(ch) in 'WF' else .57)
                    if line and width+cw>w-6:result.append(line);line='';width=0
                    line+=ch;width+=cw
                result.append(line)
            return result
        original=text
        while size>13 and len(wrap(original,size))*size*1.18>h:
            size-=.5
        text='\n'.join(wrap(original,size))
        self.items.append(('text',x,y,w,h,text,size,bold,color,align))
        sh=self.slide.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100))
        tf=sh.text_frame; tf.clear(); tf.word_wrap=False
        tf.margin_left=tf.margin_right=Inches(.02); tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
            p.font.name=FONT;p.font.size=Pt(size*.72);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.alignment=PP_ALIGN.CENTER if align=='center' else PP_ALIGN.LEFT
            p.space_before=p.space_after=Pt(0)
        return sh
    def box(self,x,y,w,h,title,body='',fill=PALE,color=BLUE):
        if body in [PALE,GREEN,GOLD,GRAY,BLUE,TEAL,AMBER]:
            raise ValueError(f'Color supplied as text body: {title}')
        self.items.append(('rect',x,y,w,h,fill,color,1.5))
        sh=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100))
        sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(1)
        if body:
            self.txt(x+7,y+5,w-14,31,title,20,True,color)
            self.txt(x+8,y+37,w-16,h-42,body,17,False,INK)
        else:self.txt(x+6,y+4,w-12,h-8,title,19,True,color)
    def panel(self,x,y,w,h,title):
        self.box(x,y,w,h,'',fill='FFFFFF',color=LINE)
        self.txt(x+8,y+4,w-16,31,title,20,True,BLUE,'left')
    def line(self,pts,arrow=True,color=BLUE,dash=False):
        self.items.append(('line',pts,arrow,color,dash))
        for i,(a,b) in enumerate(zip(pts,pts[1:])):
            s=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]/100),Inches(a[1]/100),Inches(b[0]/100),Inches(b[1]/100))
            s.line.color.rgb=RGBColor.from_string(color);s.line.width=Pt(1.1)
            ln=s._element.spPr.get_or_add_ln()
            if dash:
                d=OxmlElement('a:prstDash');d.set('val','dash');ln.append(d)
            if arrow and i==len(pts)-2:
                e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
    def export(self):
        parts=['<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="600" viewBox="0 0 1200 600">','<rect width="1200" height="600" fill="white"/>']
        for it in self.items:
            if it[0]=='rect':
                _,x,y,w,h,f,c,l=it;parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#{f}" stroke="#{c}" stroke-width="{l}"/>')
            elif it[0]=='text':
                _,x,y,w,h,t,s,b,c,a=it; lines=t.split('\n');lh=s*1.18; yy=y+(h-lh*len(lines))/2+s*.9
                for j,l in enumerate(lines):
                    anchor='middle' if a=='center' else 'start'; xx=x+w/2 if a=='center' else x
                    parts.append(f'<text x="{xx}" y="{yy+j*lh}" text-anchor="{anchor}" font-family="{FONT}" font-size="{s}" font-weight="{700 if b else 400}" fill="#{c}">{html.escape(l)}</text>')
            else:
                _,pts,arr,c,d=it; dash=' stroke-dasharray="7 5"' if d else ''
                parts.append(f'<polyline points="{" ".join(f"{x},{y}" for x,y in pts)}" fill="none" stroke="#{c}" stroke-width="1.7"{dash}/>')
                if arr:
                    (x0,y0),(x,y)=pts[-2:];ang=math.atan2(y-y0,x-x0)
                    pp=[(x,y),(x-9*math.cos(ang-.38),y-9*math.sin(ang-.38)),(x-9*math.cos(ang+.38),y-9*math.sin(ang+.38))]
                    parts.append(f'<polygon points="{" ".join(f"{a},{b}" for a,b in pp)}" fill="#{c}"/>')
        parts.append('</svg>');svg='\n'.join(parts)
        (OUT/f'{self.name}.svg').write_text(svg)
        cairosvg.svg2pdf(bytestring=svg.encode(),write_to=str(OUT/f'{self.name}.pdf'))
        cairosvg.svg2png(bytestring=svg.encode(),write_to=str(OUT/f'{self.name}.png'),output_width=1600,output_height=800)

f=Fig('01-overview','整体框架：人的方向、团队工作与真实结果')
f.box(35,105,205,155,'人类目标','方向与偏好\n关键约束\n成果体验')
f.box(305,105,245,155,'语义与工作状态','目标／规范／设计\n任务／职责／依赖\n内容／版本／证据')
f.box(615,105,245,155,'自主开发组织','委派与上下文\n模块实现与协作\n验证与局部交付')
f.box(925,105,240,155,'实际软件与使用','功能与运行状态\n用户任务与故障\n变化后的真实结果')
for a,b in [(240,305),(550,615),(860,925)]:f.line([(a,183),(b,183)])
f.box(305,350,555,150,'Workbench：理解与控制共处','按当前问题选择表达 → 体验／探索／比较 → 对象上的操作与反馈',GREEN,TEAL)
f.line([(1045,260),(1045,423),(860,423)],color=TEAL)
f.line([(305,423),(135,423),(135,260)],color=TEAL)
f.line([(550,350),(550,300),(427,300),(427,260)],color=TEAL,dash=True)
f.txt(870,330,290,35,'观察与验证证据',17,color=TEAL)

f=Fig('02-amplification','能力放大：组合、积累、反馈与交互')
for i,(t,b) in enumerate([('组合','模块边界\n协作协议\n按需委派'),('积累','设计理由\n故障教训\n可复用能力'),('反馈','定向检查\n真实使用\n交付后观察'),('交互','选择重点\n可操作表达\n准确的人类反馈')]):
    f.box(35+i*290,100,260,170,t,b,PALE if i%2==0 else GREEN,BLUE if i%2==0 else TEAL)
    f.line([(165+i*290,270),(165+i*290,325)])
f.box(35,325,1130,85,'工程能力：可承担的变化复杂度、质量与持续维护范围',fill=GREEN,color=TEAL)
f.box(35,455,350,70,'人：理解／决策／恢复','',GOLD,AMBER)
f.box(425,455,350,70,'机器：推理／工具／等待','',GRAY,INK)
f.box(815,455,350,70,'组织：协调／知识维护','',GRAY,INK)
f.txt(35,415,1130,30,'收益需扣除以下成本；增加成员或 token 不自动产生放大',17,color=INK)

f=Fig('03-semantics','语义约束：从目标到可检查的结果')
steps=[('目标','取消后可确认任务停止'),('场景','运行中取消／重复取消'),('契约','请求已收／停止已证实'),('实现','前端／调度／执行器'),('证据','状态、资源、用户体验')]
for i,(t,b) in enumerate(steps):
    f.box(35+230*i,105,210,130,t,b)
    if i<4:f.line([(245+230*i,170),(265+230*i,170)])
f.panel(35,295,540,235,'遇到冲突时追查假设')
f.box(55,350,150,75,'规范意图');f.box(230,350,150,75,'当前行为');f.box(405,350,150,75,'测试断言')
f.txt(55,445,500,65,'依据当前认可目标和具体证据判断\n修改其中过时或错误的部分',19)
f.panel(610,295,555,235,'每次变化保存足够的联系')
f.txt(630,350,515,150,'对象：改的是哪个行为或契约\n依据：目标、实例与适用条件\n状态：拟议／实现／已验证\n影响：责任模块及需更新的检查',20,False,INK,'left')

f=Fig('04-organization','团队结构与软件结构：相关但分别建模')
f.panel(35,90,1130,190,'组织关系：长期职责与临时协调')
f.box(65,145,205,95,'Leader','目标／偏好／承诺')
f.box(330,145,230,95,'任务负责人','跨模块方案与交付')
f.box(620,145,230,95,'模块成员 A','执行器长期经验')
f.box(910,145,230,95,'模块成员 B','任务状态长期经验')
f.line([(270,193),(330,193)]);f.line([(560,193),(620,193)])
f.line([(445,240),(445,260),(1025,260),(1025,240)])
f.panel(35,350,1130,185,'软件关系：服务职责与运行协议')
f.box(65,410,220,85,'用户界面','操作与反馈')
f.box(390,410,290,85,'任务状态服务','请求记录与状态查询')
f.box(840,410,290,85,'执行器','实际停止与资源释放')
f.line([(285,453),(390,453)]);f.line([(680,453),(840,453)])
f.line([(735,240),(735,320),(985,320),(985,410)],False,TEAL,True)
f.line([(1025,240),(1155,240),(1155,330),(535,330),(535,410)],False,TEAL,True)
f.txt(85,294,430,32,'虚线关联：谁维护什么；一项工作可跨多个模块',16,color=TEAL)

f=Fig('05-context','上下文与经验：选择、使用、提炼与恢复')
f.box(35,105,270,175,'持久材料','规范与模块契约\n历史判断与证据\n未完任务与入口')
f.box(395,105,330,175,'当前任务上下文','目标＋必要依赖\n相关经验＋当前差异\n工具结果的局部证据')
f.box(815,105,350,175,'执行与核对','采取行动／观察结果\n解释失败／修正假设\n完成或交接未完工作')
f.line([(305,193),(395,193)]);f.line([(725,193),(815,193)])
f.box(395,365,330,155,'整理后的经验','适用条件／反例\n可复用步骤／核对方法\n关联原始证据',GREEN,TEAL)
f.line([(990,280),(990,442),(725,442)],color=TEAL)
f.line([(395,442),(170,442),(170,280)],color=TEAL,dash=True)
f.txt(60,305,285,40,'按任务检索，避免全部注入',17)
f.txt(805,355,360,85,'恢复时保留：\n职责、承诺、关键决定、工作入口',18)

f=Fig('06-delivery','并行变化中的小批量交付')
rows=[('模块 A',120),('模块 B',235),('发布工具',350),('使用与验证',465)]
for t,y in rows:
    f.txt(30,y,125,60,t,20,True);f.line([(165,y+30),(1165,y+30)],False,LINE)
f.box(185,115,240,65,'实现取消请求');f.box(560,115,250,65,'继续完善输出保留')
f.box(300,230,235,65,'执行器响应与状态');f.box(850,230,275,65,'下一项独立变化')
f.box(560,335,180,85,'捕获实际内容','范围＋依赖');f.box(780,335,175,85,'合成与切换','保留并行成果');f.box(995,335,170,85,'实际交付记录','内容＋结果')
f.line([(425,147),(485,147),(485,377),(560,377)])
f.line([(535,262),(545,262),(545,377),(560,377)])
f.line([(740,377),(780,377)]);f.line([(955,377),(995,377)])
f.box(590,460,260,75,'交付后观察与用户体验','',GREEN,TEAL)
f.box(910,460,255,75,'故障 → 局部修复','',GOLD,AMBER)
f.line([(1080,420),(1080,445),(720,445),(720,460)],color=TEAL)
f.line([(850,497),(910,497)],color=TEAL)
f.txt(185,82,950,30,'时间向右；仅示意依赖顺序，未表示真实耗时',16,color=INK)

f=Fig('07-evidence','验证证据的局部有效性')
cols=[35,285,510,735,960]; widths=[230,205,205,205,205]
heads=['检查对象','任务服务','执行器','当前变化','下一步']
for x,w,t in zip(cols,widths,heads):f.box(x,95,w,60,t)
data=[['请求幂等性','已验证','无关','未涉及','复用证据'],['停止确认','依赖执行器','已验证','执行器修改','重查受影响项'],['用户取消旅程','已验证','已验证','多模块组合','运行定向 TAU'],['无关导出任务','已验证','无关','未涉及','保留原结论']]
for j,row in enumerate(data):
    for i,t in enumerate(row):f.box(cols[i],180+j*78,widths[i],62,t,fill=GREEN if i==4 and j in (0,3) else (GOLD if i==4 else 'FFFFFF'),color=TEAL if i==4 and j in (0,3) else INK)
f.txt(40,505,1115,40,'证据 = 任务／契约 + 实际内容与环境 + 观察结果 + 适用条件',21,True)

f=Fig('08-validation','自适应验证：决定投入并完成修复')
for x,t,b in [(35,'变化与故障','相关模块／契约\n已有检查与未覆盖项'),(430,'判断验证范围','预期信息价值\n运行成本与影响范围'),(825,'执行选定检查','代码测试／任务体验\n环境与结果记录')]:f.box(x,105,340,145,t,b)
f.line([(375,178),(430,178)]);f.line([(770,178),(825,178)])
f.box(35,365,340,145,'更新证据与下一步','保留有效范围\n定位剩余未知',GREEN,TEAL)
f.box(430,365,340,145,'修复与必要复查','产品／测试／环境\n相关责任成员',GOLD,AMBER)
f.box(825,365,340,145,'解释结果','通过：范围内成立\n失败：区分问题来源')
f.line([(995,250),(995,365)]);f.line([(825,437),(770,437)]);f.line([(430,437),(375,437)])
f.line([(205,365),(205,305),(600,305),(600,250)],color=TEAL,dash=True)
f.txt(450,275,655,30,'按新证据调整；无需为其他无关变化反复重跑',17,color=TEAL)

f=Fig('09-projection','工作投影：真实对象、生成表达与操作回传')
f.box(35,110,300,390,'工作语义与状态','目标／模块／协议\n任务／依赖／责任\n版本／证据／结果\n\n稳定对象标识\n各状态的实际来源')
f.box(435,110,300,160,'投影生成过程','人的当前问题与熟悉程度\n选择信息、粒度与媒体\n模板与表达经验')
f.box(835,110,330,160,'可理解的界面','结构／演示／比较\n重点＋逐步深入\n来源与状态可查看',GREEN,TEAL)
f.box(435,355,300,145,'动作解析与结果','对象＋参数＋预期作用\n执行后读取真实状态')
f.box(835,355,330,145,'人的直接参与','体验、标注、选择\n提出新的约束与问题',GREEN,TEAL)
f.line([(335,190),(435,190)]);f.line([(735,190),(835,190)])
f.line([(1000,270),(1000,355)],color=TEAL);f.line([(835,427),(735,427)],color=TEAL)
f.line([(435,427),(335,427)],color=TEAL)
f.line([(585,355),(585,270)],color=TEAL,dash=True)

f=Fig('10-architecture','架构理解指南：层次展开与协议探索')
f.panel(35,90,690,445,'系统 > 任务服务 > 取消协议')
f.box(55,145,650,100,'顶层：任务平台','界面、任务服务、执行器、结果存储')
f.line([(380,245),(380,275)])
f.box(55,280,195,95,'请求处理','权限／幂等');f.box(280,280,195,95,'状态协调','接收／停止确认');f.box(510,280,195,95,'执行适配','实际停止')
f.line([(250,327),(280,327)]);f.line([(475,327),(510,327)])
f.box(55,420,650,80,'关联改进：请求接收与实际停止分别呈现','',GOLD,AMBER)
f.panel(770,90,395,445,'点击协议：用具体过程解释')
f.box(790,150,155,60,'任务服务');f.box(985,150,155,60,'执行器')
f.line([(867,220),(867,495)],False,LINE);f.line([(1062,220),(1062,495)],False,LINE)
for y,t,rev in [(260,'停止请求',False),(330,'请求已接收',True),(425,'停止已确认',True)]:
    f.line([(1062 if rev else 867,y),(867 if rev else 1062,y)],color=TEAL if rev else BLUE)
    f.txt(865,y-30,205,28,t,16)

f=Fig('11-workbench','Workbench 示例：体验结果、解释机制、直接反馈')
f.box(35,90,1130,55,'任务平台  /  运行中的任务  /  取消体验','',GRAY,INK)
f.panel(35,165,220,370,'导航与关注点')
f.txt(55,215,180,260,'功能体验\n\n架构与协议\n\n本次变化\n\n验证范围',21,False,INK,'left')
f.panel(280,165,530,370,'任务 A：模拟演示')
f.box(305,220,480,85,'运行中 → 请求已接收 → 停止已确认')
f.box(305,340,140,55,'启动任务');f.box(475,340,140,55,'请求取消');f.box(645,340,140,55,'重置演示')
f.txt(305,425,480,75,'观察：资源释放完成后确认停止\n可保留的输出在结果区继续查看',19)
f.panel(835,165,330,370,'机制与证据')
f.txt(855,215,290,155,'协议：取消请求／停止确认\n证据：关联的测试与记录\n范围：正常／重复／并发取消',17,False,INK,'left')
f.box(855,395,290,95,'在此步骤反馈','“接收请求后应显示等待状态”',GREEN,TEAL)

f=Fig('12-guides','六类表达指南：按人的理解任务选择')
cells=[('系统结构','全貌 → 模块 → 协议','标注职责／查看影响'),('功能体验','任务 → 操作 → 结果','尝试行为／反馈步骤'),('运行机制','状态 → 事件 → 分支','调整条件／暂停观察'),('方案比较','同一场景 → 候选差异','补充约束／选择方案'),('阶段进展','近期变化 → 成果与阻塞','体验成果／调整关注'),('未来演化','当前系统 → 拟议路径','调整优先级／查看依赖')]
for i,(t,b,c) in enumerate(cells):
    x=35+(i%3)*390;y=100+(i//3)*230
    f.box(x,y,350,185,t,b+'\n\n'+c,GREEN if i%2 else PALE,TEAL if i%2 else BLUE)

f=Fig('13-case-sequence','构造案例：取消请求与完成事件的竞争')
xs=[135,435,735,1065];titles=['用户界面','任务状态服务','执行器','资源／结果']
for x,t in zip(xs,titles):f.box(x-100,90,200,65,t);f.line([(x,160),(x,530)],False,LINE)
seq=[(205,0,1,'提交取消请求'),(265,1,0,'请求已接收'),(325,1,2,'转发停止意图'),(385,2,3,'释放资源／保留输出'),(445,2,1,'终态确认与版本'),(505,1,0,'显示实际终态')]
for y,a,b,t in seq:
    f.line([(xs[a],y),(xs[b],y)],color=TEAL if b<a else BLUE)
    f.txt(min(xs[a],xs[b])+10,y-28,abs(xs[b]-xs[a])-20,27,t,16)
f.txt(745,190,395,95,'并发分支：完成事件可能先到达\n状态服务核对实际终态与版本\n界面据此说明结果',18,False,AMBER)

f=Fig('14-case-iteration','完整工作示例：从目标进入下一轮变化')
for i,(t,b) in enumerate([('意图','用户需要确认任务停止'),('分工','状态／执行／界面职责'),('实现','契约＋局部修改'),('交付','实际内容与定向证据')]):
    f.box(35+i*290,105,260,145,t,b)
    if i<3:f.line([(295+i*290,178),(325+i*290,178)])
f.box(905,355,260,155,'TAU 发现问题','等待状态缺少说明\n刷新后才看到终态',GOLD,AMBER)
f.box(615,355,260,155,'回到相关职责','补充等待与状态更新\n复查体验路径',GOLD,AMBER)
f.box(325,355,260,155,'展示与反馈','人体验新的状态演示\n提出保留中间结果',GREEN,TEAL)
f.box(35,355,260,155,'下一轮意图','更新输出保留契约\n复用有效验证',GREEN,TEAL)
f.line([(1035,250),(1035,355)])
for a,b in [(905,875),(615,585),(325,295)]:f.line([(a,432),(b,432)],color=TEAL)
f.line([(165,355),(165,250)],color=TEAL,dash=True)

f=Fig('15-evaluation','评价设计：多轮变化、对照与独立核对')
f.panel(35,95,340,440,'比较条件')
f.txt(55,150,300,315,'单 Agent 基线\n模块组织\n组织＋持续经验\n组织＋交互投影\n\n相同初始代码与任务\n记录模型、预算、工具',21,False,INK,'left')
f.panel(420,95,340,440,'连续任务序列')
for j,t in enumerate(['新增能力','跨模块需求变化','故障诊断与修复','后续维护与交接']):
    f.box(445,150+j*85,290,62,t)
    if j<3:f.line([(590,212+j*85),(590,235+j*85)])
f.panel(805,95,360,440,'共同评价维度')
f.txt(830,150,315,315,'任务结果与使用质量\n修改的传播与返工\n交付及修复时间\n人类理解与决策时间\n模型和工具资源\n\n固定检查＋人类任务\n记录失败与未完成项',20,False,INK,'left')
f.line([(375,315),(420,315)]);f.line([(760,315),(805,315)])

f=Fig('16-cost','表达投入的价值：资源分配与可观察收益')
f.box(35,105,315,165,'资源投入','代码实现与修复\n验证、协调与记忆\n面向人的表达与交互')
f.box(440,105,315,165,'可能产生的收益','更快建立理解\n更准确表达偏好\n更早发现方向偏差',GREEN,TEAL)
f.box(845,105,320,165,'需要共同观察','人的时间与任务正确性\n后续误解与返工\n总费用与交付时间')
f.line([(350,188),(440,188)]);f.line([(755,188),(845,188)])
f.panel(35,340,1130,190,'研究假设与判定条件')
f.txt(65,395,1070,100,'展示与交互的 token 消耗可能接近代码生成；比例需随任务观察。\n投入增加是否值得，取决于人的理解、控制质量与后续工程结果。\n复用模板、局部更新与按需生成可改变成本，不设置固定预算比例。',22)

for f in figures:f.export()
prs.save(ROOT/'figures-editable.pptx')
manifest=[{'index':i+1,'name':f.name,'native_shapes':len(f.slide.shapes),'scene_objects':len(f.items)} for i,f in enumerate(figures)]
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'figures':len(figures),'native_shapes':sum(m['native_shapes'] for m in manifest),'pptx':str(ROOT/'figures-editable.pptx')},ensure_ascii=False))
