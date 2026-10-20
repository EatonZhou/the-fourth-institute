"""Check citations, figure exports, PDF diagnostics and native PPT editability."""
from pathlib import Path
import re, json, subprocess, zipfile, xml.etree.ElementTree as ET
from pptx import Presentation
from PIL import Image, ImageOps, ImageDraw
root=Path(__file__).resolve().parents[1]
build=root/'build'; build.mkdir(exist_ok=True)
tex='\n'.join(p.read_text() for p in [root/'main.tex',*sorted((root/'chapters').glob('*.tex'))])
refs=(root/'references.tex').read_text()
bib=set(re.findall(r'\\bibitem\{([^}]+)\}',refs))
cited={k.strip() for group in re.findall(r'\\cite\{([^}]+)\}',tex) for k in group.split(',')}
labels=re.findall(r'\\label\{([^}]+)\}',tex)
figs=re.findall(r'\\fig\{([^}]+)\}\{.*?\}\{([^}]+)\}',tex,re.S)
labels += ['fig:'+label for _,label in figs]
usedrefs=set(re.findall(r'\\ref\{([^}]+)\}',tex))
log=(build/'main.log').read_text()
errors=[x for x in log.splitlines() if any(y in x for y in ['Missing character:','Overfull','undefined','LaTeX Error:'])]
info=subprocess.run(['pdfinfo',str(root/'article.pdf')],capture_output=True,text=True,check=True).stdout
page_count=int(re.search(r'Pages:\s+(\d+)',info)[1])
txt=subprocess.run(['pdftotext','-layout',str(root/'article.pdf'),'-'],capture_output=True,text=True,check=True).stdout
(build/'article-extracted.txt').write_text(txt)
pages=txt.split('\f')[:page_count]
deck=Presentation(root/'figures-editable.pptx')
shape_count=0; pictures=0;outside=[]
for i,slide in enumerate(deck.slides,1):
    for shape in slide.shapes:
        shape_count+=1
        if shape.shape_type==13:pictures+=1
        if shape.left<0 or shape.top<0 or shape.left+shape.width>deck.slide_width+10 or shape.top+shape.height>deck.slide_height+10:outside.append([i,shape.name])
report={'pages':page_count,'chapters':len(list((root/'chapters').glob('*.tex'))),'chinese_characters':len(re.findall(r'[\u4e00-\u9fff]',tex)),
        'figures':len(figs),'references':len(bib),'uncited_references':sorted(bib-cited),'missing_citations':sorted(cited-bib),
        'undefined_refs':sorted(usedrefs-set(labels)),'duplicate_labels':sorted({x for x in labels if labels.count(x)>1}),
        'missing_figure_exports':[f for f,_ in figs if not (root/'figures'/f'{f}.pdf').is_file()],
        'latex_errors_or_overflows':errors,'ppt_slides':len(deck.slides),'ppt_native_shapes':shape_count,'ppt_bitmap_pictures':pictures,'ppt_out_of_bounds':outside,
        'page_text_characters':[len(p.strip()) for p in pages]}
(build/'checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k!='page_text_characters'},ensure_ascii=False,indent=2))
if any(report[k] for k in ['missing_citations','undefined_refs','duplicate_labels','missing_figure_exports','latex_errors_or_overflows','ppt_out_of_bounds']):raise SystemExit(1)
thumbs=build/'page-previews';thumbs.mkdir(exist_ok=True)
subprocess.run(['pdftoppm','-scale-to','500','-png',str(root/'article.pdf'),str(thumbs/'page')],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
pngs=sorted(thumbs.glob('page-*.png'))
for batch in range((len(pngs)+11)//12):
    chosen=pngs[batch*12:(batch+1)*12];sheet=Image.new('RGB',(4*380,3*550),'#d6dce1');d=ImageDraw.Draw(sheet)
    for j,path in enumerate(chosen):
        im=Image.open(path).convert('RGB');im.thumbnail((360,510));x=(j%4)*380+10;y=(j//4)*550+28
        sheet.paste(im,(x,y));d.text((x,y-20),f'Page {batch*12+j+1}',fill='black')
    sheet.save(build/f'pages-contact-{batch+1}.png')
figpngs=sorted((root/'figures').glob('*.png'))
for batch in range((len(figpngs)+3)//4):
    sheet=Image.new('RGB',(1600,860),'#d6dce1');d=ImageDraw.Draw(sheet)
    for j,path in enumerate(figpngs[batch*4:(batch+1)*4]):
        im=Image.open(path).convert('RGB');im.thumbnail((780,390));x=(j%2)*800+10;y=(j//2)*430+27
        sheet.paste(im,(x,y));d.text((x,y-20),path.stem,fill='black')
    sheet.save(build/f'figures-contact-{batch+1}.png')
